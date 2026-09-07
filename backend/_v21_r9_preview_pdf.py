"""V2.1.0 R9：pm_template v1.2 三端一致性 + 真实 PDF 渲染/下载验证脚本。

零 API Key：不调 LLM/Embedding/外发请求；所有输出写入进程级临时 RESUME_DATA_DIR。
直接驱动 TemplateRenderer（docx）/ pdf_renderer（PDF）/ _build_doc_preview（预览 JSON），
并用 mini FastAPI + TestClient 走真实 download 路由做正反向验证。

运行：python _v21_r9_preview_pdf.py
要求：exit 0 且末行 PASS=<N> FAIL=0。
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

# ── 进程级临时数据根：必须在任何 backend import 之前设置 ──
_R9_TMP_ROOT = tempfile.mkdtemp(prefix="v21_r9_")
os.environ["RESUME_DATA_DIR"] = _R9_TMP_ROOT

BACKEND_ROOT = Path(__file__).resolve().parent
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

PASS_COUNT = 0
FAILURES: list[str] = []


def _assert(cond: bool, label: str, detail: str = "") -> None:
    global PASS_COUNT
    if cond:
        PASS_COUNT += 1
        print(f"  [PASS] {label}")
    else:
        msg = f"  [FAIL] {label}" + (f" — {detail}" if detail else "")
        print(msg, file=sys.stderr)
        FAILURES.append(label)


def _section(title: str) -> None:
    print(f"\n[{title}]")


def _norm(s: str) -> str:
    """去除全部空白（含换行）——用于跨 docx/pdf 文本对照。"""
    return "".join(s.split())


# ═══════════════════════════════════════════════════════════════════════════
# 固定虚构 fixture（确定性中文内容，避开敏感真实信息）
# ═══════════════════════════════════════════════════════════════════════════
def _make_resume_doc():
    from models.resume_document import (
        EducationItem, Profile, ProjectItem, ResumeDocument, SkillGroup, WorkItem,
    )
    return ResumeDocument(
        profile=Profile(
            name="林晓芸",
            phone="139-1234-5678",
            email="linxiaoyun@example.com",
            location="上海",
            target_position="高级产品经理",
            summary=None,
        ),
        education=[
            EducationItem(
                school="云帆理工大学",
                major="产品设计",
                degree="本科",
                start_time="2020.09",
                end_time="2024.06",
                description="主修课程：用户体验、交互设计、用户研究、数据分析；GPA 3.8/4.0",
                experience_id="exp-edu-1",
            ),
        ],
        work=[
            WorkItem(
                company="云帆科技",
                role="产品实习生",
                start_time="2024.07",
                end_time="2024.12",
                bullets=[
                    "完成 6 场用户访谈与 200 份问卷分析，输出洞察报告支撑改版决策",
                    "协同研发上线 2 个迭代版本，跟踪 18 个需求从评审到验收全流程",
                ],
                experience_id="exp-work-1",
            ),
        ],
        projects=[
            ProjectItem(
                name="智能日程助手 App",
                role="产品负责人",
                start_time="2025.01",
                end_time="2025.06",
                bullets=[
                    "定义核心场景与功能优先级，产出 PRD 与可交互原型并完成 3 轮可用性测试",
                ],
                experience_id="exp-proj-1",
            ),
        ],
        skills=[
            SkillGroup(category="产品技能", items=["需求分析", "原型设计", "PRD 撰写", "A/B 测试"]),
            SkillGroup(category="工具", items=["Figma", "Axure", "SQL"]),
        ],
        awards=[],  # 空可选章节：三端必须一致隐藏
    )


def _all_bullets(doc) -> list[str]:
    """fixture 中所有应出现的 bullet 文本（教育 desc 计 1，work/project bullets）。"""
    out: list[str] = []
    edu = doc.education[0]
    if edu.description:
        out.append(edu.description)
    out.extend(doc.work[0].bullets)
    out.extend(doc.projects[0].bullets)
    return out


def _read_docx_text(path) -> list[str]:
    from docx import Document
    doc = Document(path)
    return [p.text for p in doc.paragraphs if (p.text or "").strip()]


def _read_pdf_text(path) -> str:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join((pg.extract_text() or "") for pg in pdf.pages)


def _pdf_lines(path) -> list[str]:
    txt = _read_pdf_text(path)
    return [ln for ln in txt.split("\n") if ln.strip()]


def _docx_paragraphs(path) -> list[str]:
    return _read_docx_text(path)


def _visible_chapter_keys(lines) -> list[str]:
    """按出现顺序识别章节锚点 → 规范化 key（education/work/project/skills/awards/summary）。"""
    anchor = {
        "教育背景": "education",
        "实习经历": "work",
        "项目经历": "project",
        "技能专长": "skills",
        "荣誉奖项": "awards",
        "自我评价": "summary",
    }
    keys: list[str] = []
    for ln in lines:
        n = _norm(ln)
        for title, key in anchor.items():
            if n == _norm(title):
                if key not in keys:
                    keys.append(key)
                break
    return keys


def _preview_sections(doc):
    from services.resume_generation_service import _build_doc_preview
    return _build_doc_preview(doc)


def _count_docx_bullets(path) -> int:
    return sum(1 for ln in _read_docx_text(path) if _norm(ln).startswith("⚫"))


def _count_pdf_bullets(path) -> int:
    return sum(1 for ln in _pdf_lines(path) if _norm(ln).startswith("●"))


# ═══════════════════════════════════════════════════════════════════════════
# 1) 三端产物生成（docx/pdf/preview）
# ═══════════════════════════════════════════════════════════════════════════
def _produce_all():
    """生成并落盘 docx + pdf；返回 (doc, docx_path, pdf_bytes, pdf_path, warnings)。"""
    from core.config import settings
    from services.template_renderer import TemplateRenderer
    from services import pdf_renderer

    out_dir = settings.DOCX_OUTPUT_DIR
    os.makedirs(out_dir, exist_ok=True)
    user_id = "r9user"
    template_id = "pm_template"

    doc = _make_resume_doc()

    # Word：与主链路同源（TemplateRenderer + docx_writer）
    renderer = TemplateRenderer(template_id, backend_root=str(BACKEND_ROOT))
    docx_doc, docx_warnings, _stats = renderer.render(doc)
    docx_name = f"resume_{user_id}_{template_id}.docx"
    docx_path = os.path.join(out_dir, docx_name)
    docx_doc.save(docx_path)

    # PDF：同一 resume_doc → 真实 reportlab 产物
    pdf_bytes, pdf_warnings = pdf_renderer.render(doc, template_id, str(BACKEND_ROOT))
    pdf_name = f"resume_{user_id}_{template_id}.pdf"
    pdf_path = os.path.join(out_dir, pdf_name)
    with open(pdf_path, "wb") as f:
        f.write(pdf_bytes)

    return doc, docx_path, pdf_bytes, pdf_path, (docx_warnings, pdf_warnings)


# ═══════════════════════════════════════════════════════════════════════════
# 2) PDF 真实性与文件名 / 单页
# ═══════════════════════════════════════════════════════════════════════════
def test_pdf_authentic(pdf_bytes, pdf_path, template_id) -> None:
    _section("1) PDF 真实性 / 文件名 / 页数 / 分隔线")
    _assert(pdf_bytes[:5] == b"%PDF-", "PDF 产物以 %PDF 头开始", str(pdf_bytes[:8]))
    _assert(len(pdf_bytes) > 800, "PDF 产物非空（字节数足够）", f"len={len(pdf_bytes)}")
    _assert(
        os.path.basename(pdf_path) == f"resume_r9user_{template_id}.pdf",
        "PDF 文件名符合 resume_<user>_<template>.pdf",
        os.path.basename(pdf_path),
    )
    import pdfplumber
    with pdfplumber.open(pdf_path) as pdf:
        _assert(len(pdf.pages) == 1, "PDF 恰好 1 页", f"pages={len(pdf.pages)}")
        page = pdf.pages[0]
        horiz = [l for l in (page.lines or []) if abs(l["top"] - l["bottom"]) < 0.5 and l["width"] > 100]
        _assert(len(horiz) >= 4, "PDF 含 ≥4 条章节标题底分隔线（4 个可见章节）", f"n={len(horiz)}")
        _assert(len(page.rects or []) >= 1, "PDF 含照片占位框矩形（rect）", f"n={len(page.rects or [])}")
        left_margin_pt = 1.13 * 28.3465
        if page.chars:
            min_x = min(c["x0"] for c in page.chars)
            _assert(
                abs(min_x - left_margin_pt) < 6.0,
                "PDF 文本左缘≈模板左页边距 1.13cm",
                f"min_x={min_x:.1f}pt expect≈{left_margin_pt:.1f}pt",
            )
        else:
            _assert(False, "PDF 有可提取文本（chars 非空）")


# ═══════════════════════════════════════════════════════════════════════════
# 3) 三端内容一致性
# ═══════════════════════════════════════════════════════════════════════════
def test_triplet_consistency(doc, docx_path, pdf_path) -> None:
    _section("2) 三端一致性：profile / 章节顺序 / 条目 / bullets / 空章节")

    docx_text = " ".join(_read_docx_text(docx_path))
    docx_n = _norm(docx_text)
    pdf_n = _norm(_read_pdf_text(pdf_path))
    preview = _preview_sections(doc)
    pv_by_key = {s.section: s for s in preview}

    prof = doc.profile
    # profile：姓名/联系方式/目标岗位
    pv_personal = pv_by_key["personal"].entries[0]
    _assert(
        _norm(prof.name) in docx_n and _norm(prof.name) in pdf_n
        and pv_personal.heading == prof.name,
        "姓名在 docx/pdf/preview 三端一致",
    )
    contact_vals = [prof.phone, prof.email, prof.location, prof.target_position]
    _assert(
        all(_norm(v) in docx_n for v in contact_vals)
        and all(_norm(v) in pdf_n for v in contact_vals)
        and _norm(prof.phone) in _norm(pv_personal.subhead or "")
        and _norm(prof.email) in _norm(pv_personal.subhead or "")
        and any(prof.target_position in b for b in pv_personal.bullets),
        "电话/邮箱/所在地/目标岗位在 docx/pdf/preview 三端一致",
    )

    # 章节顺序一致：docx/pdf 可见章节 key 相同，且 preview 非 personal 部分 = 同序列
    docx_keys = _visible_chapter_keys(_read_docx_text(docx_path))
    pdf_keys = _visible_chapter_keys(_pdf_lines(pdf_path))
    expected_visible = ["education", "work", "project", "skills"]
    _assert(docx_keys == expected_visible, "docx 可见章节顺序 == 模板顺序 education/work/project/skills", f"{docx_keys}")
    _assert(pdf_keys == expected_visible, "pdf 可见章节顺序 == 模板顺序 education/work/project/skills", f"{pdf_keys}")
    pv_visible = [s.section for s in preview if s.section != "personal"]
    _assert(pv_visible == expected_visible, "preview JSON 可见章节顺序与 Word/PDF 一致", f"{pv_visible}")
    _assert([s.section for s in preview][0] == "personal", "preview personal 位于首位")

    # 空章节处理一致：awards/summary 为空 → 三端均隐藏
    _assert(
        "荣誉奖项" not in docx_n and "自我评价" not in docx_n
        and "荣誉奖项" not in pdf_n and "自我评价" not in pdf_n
        and "awards" not in pv_by_key and "summary" not in pv_by_key,
        "空可选章节（荣誉奖项/自我评价）docx/pdf/preview 一致隐藏",
    )

    # education 条目：标题字段（school/major/degree/时间）在 Word 与 PDF 均有呈现
    edu = doc.education[0]
    pv_edu = pv_by_key["education"].entries[0]
    _assert(
        _norm(edu.school) in docx_n and _norm(edu.school) in pdf_n
        and edu.school in pv_edu.heading,
        "education.school 三端一致",
    )
    _assert(
        _norm(edu.major) in docx_n and _norm(edu.major) in pdf_n
        and _norm(edu.major) in _norm(pv_edu.heading),
        "education.major 三端一致",
    )
    _assert(
        _norm(edu.degree) in docx_n and _norm(edu.degree) in pdf_n,
        "education.degree 在 docx/pdf 呈现（Word 模板 专业（学历）行布局）",
    )
    _assert(
        _norm(edu.start_time) in docx_n and _norm(edu.start_time) in pdf_n
        and _norm(edu.start_time) in _norm(pv_edu.subhead or ""),
        "education 时间（起）三端一致",
    )
    _assert(
        _norm(edu.end_time) in docx_n and _norm(edu.end_time) in pdf_n
        and _norm(edu.end_time) in _norm(pv_edu.subhead or ""),
        "education 时间（止）三端一致",
    )

    # 全量 bullet 文本逐条一致：每一条都出现在 docx/pdf/preview
    bullets_ok = True
    detail = ""
    for b in _all_bullets(doc):
        nb = _norm(b)
        if nb not in docx_n:
            bullets_ok = False
            detail = f"docx 缺 bullet: {b!r}"
            break
        if nb not in pdf_n:
            bullets_ok = False
            detail = f"pdf 缺 bullet: {b!r}"
            break
    # preview bullets（education 的 description 也进 bullets）
    pv_bullets_flat = []
    for sec in preview:
        for ent in sec.entries:
            pv_bullets_flat.extend(ent.bullets or [])
    for b in _all_bullets(doc):
        nb = _norm(b)
        if nb not in [_norm(x) for x in pv_bullets_flat]:
            bullets_ok = False
            detail = f"preview 缺 bullet: {b!r}"
            break
    _assert(bullets_ok, "全部 bullet 文字在 docx/pdf/preview 逐条一致", detail)

    # bullet 数量一致（docx⚫行 == pdf●行 == preview work/project/education bullets 数）
    # 说明：preview 的 personal（联系方式行）与 skills（技能行）不是模板 bullet 行，不计入。
    expected_n = len(_all_bullets(doc))
    docx_bullets = _count_docx_bullets(docx_path)
    pdf_bullets = _count_pdf_bullets(pdf_path)
    pv_body_bullets = sum(
        len(ent.bullets or [])
        for sec in preview
        if sec.section in {"education", "work", "project"}
        for ent in sec.entries
    )
    _assert(
        docx_bullets == expected_n and pdf_bullets == expected_n and pv_body_bullets == expected_n,
        "bullet 数量 docx/pdf/preview 三端一致",
        f"docx={docx_bullets} pdf={pdf_bullets} preview={pv_body_bullets} expect={expected_n}",
    )

    # work / project 标题行：company/role/name 与时间
    wk = doc.work[0]
    pv_wk = pv_by_key["work"].entries[0]
    _assert(
        _norm(wk.company) in docx_n and _norm(wk.company) in pdf_n and wk.company in pv_wk.heading
        and _norm(wk.role) in docx_n and _norm(wk.role) in pdf_n and wk.role in pv_wk.heading
        and _norm(wk.start_time) in _norm(pv_wk.subhead or "") and _norm(wk.end_time) in _norm(pv_wk.subhead or ""),
        "work 条目标题（公司/职位/时间）三端一致",
    )
    pj = doc.projects[0]
    pv_pj = pv_by_key["project"].entries[0]
    _assert(
        _norm(pj.name) in docx_n and _norm(pj.name) in pdf_n and pj.name in pv_pj.heading
        and _norm(pj.role) in docx_n and _norm(pj.role) in pdf_n and pj.role in pv_pj.heading
        and _norm(pj.start_time) in _norm(pv_pj.subhead or "") and _norm(pj.end_time) in _norm(pv_pj.subhead or ""),
        "project 条目标题（名称/角色/时间）三端一致",
    )

    # skills：分类与技能项
    skills_ok = True
    sk_detail = ""
    for g in doc.skills:
        line_n = _norm(f"{g.category}：{'、'.join(g.items)}")
        if line_n not in docx_n or line_n not in pdf_n:
            skills_ok = False
            sk_detail = f"docx/pdf 缺技能行: {g.category}"
            break
        pv_ent = [e for e in pv_by_key["skills"].entries if e.heading == g.category]
        if not pv_ent or sorted(pv_ent[0].bullets) != sorted(g.items):
            skills_ok = False
            sk_detail = f"preview 技能条目不符: {g.category}"
            break
    _assert(skills_ok, "skills 分类与技能项三端一致", sk_detail)


# ═══════════════════════════════════════════════════════════════════════════
# 4) download 端点正反向（真实 ASGI：200 + application/pdf / 缺失 4xx）
# ═══════════════════════════════════════════════════════════════════════════
def test_download_endpoint(pdf_bytes, pdf_path, docx_path) -> None:
    _section("3) download 端点正反向")
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from api.routes import template as template_route

    app = FastAPI()
    app.include_router(template_route.router, prefix="/api/template")
    client = TestClient(app)

    pdf_name = os.path.basename(pdf_path)
    docx_name = os.path.basename(docx_path)

    # 正向：PDF 取回 200 + application/pdf + 内容一致
    r = client.get("/api/template/download", params={"path": f"output/{pdf_name}"})
    _assert(r.status_code == 200, "download pdf 正向返回 200", f"status={r.status_code}")
    _assert(
        r.headers.get("content-type", "").lower() == "application/pdf",
        "download pdf Content-Type == application/pdf",
        r.headers.get("content-type"),
    )
    _assert(r.content == pdf_bytes, "download pdf 字节与生成文件一致")

    # 正向：Word（旧链路回归）
    r2 = client.get("/api/template/download", params={"path": f"output/{docx_name}"})
    _assert(
        r2.status_code == 200
        and "wordprocessingml.document" in r2.headers.get("content-type", ""),
        "download docx 仍返回 200 + docx MIME（旧链路不破坏）",
        f"status={r2.status_code} ct={r2.headers.get('content-type')}",
    )

    # 反向：缺失文件 → 真实 4xx（404），绝不假装成功
    r3 = client.get("/api/template/download", params={"path": "output/not_exist_resume.pdf"})
    _assert(
        400 <= r3.status_code < 500,
        "download pdf 缺失文件返回真实 4xx 状态",
        f"status={r3.status_code}",
    )
    # 反向：非法路径穿越 → 400
    r4 = client.get("/api/template/download", params={"path": "../secret"})
    _assert(
        400 <= r4.status_code < 500,
        "download 非法 path（穿越）返回真实 4xx",
        f"status={r4.status_code}",
    )


# ═══════════════════════════════════════════════════════════════════════════
# 5) 契约回归：响应新字段可选 + py_compile + 模块 import
# ═══════════════════════════════════════════════════════════════════════════
def test_response_contract_and_compile() -> None:
    _section("4) 响应契约回归 + py_compile + 模块 import")
    from api.schemas import BuildCounts, BuildMeta, RenderStats, ResumeDocxGenerateResponse

    # 默认构造（新字段缺省 None，旧契约不破坏）
    resp = ResumeDocxGenerateResponse(
        file_path="output/x.docx", file_name="x.docx",
        download_url="/api/template/download?path=output/x.docx",
        build_counts=BuildCounts(), build_meta=BuildMeta(), render_stats=RenderStats(),
    )
    _assert(
        resp.pdf_file_name is None and resp.pdf_download_url is None,
        "ResumeDocxGenerateResponse 默认 pdf_file_name/pdf_download_url=None（向后兼容）",
    )
    resp2 = ResumeDocxGenerateResponse(
        file_path="output/x.docx", file_name="x.docx",
        download_url="/api/template/download?path=output/x.docx",
        build_counts=BuildCounts(), build_meta=BuildMeta(), render_stats=RenderStats(),
        pdf_file_name="resume_u_pm_template.pdf",
        pdf_download_url="/api/template/download?path=output/resume_u_pm_template.pdf",
    )
    dump = resp2.model_dump()
    _assert(
        dump["pdf_file_name"] == "resume_u_pm_template.pdf"
        and dump["pdf_download_url"].endswith(".pdf"),
        "携带 pdf_* 字段可序列化往返",
    )

    import py_compile
    targets = [
        BACKEND_ROOT / "services" / "pdf_renderer.py",
        BACKEND_ROOT / "services" / "resume_generation_service.py",
        BACKEND_ROOT / "api" / "schemas.py",
        BACKEND_ROOT / "api" / "routes" / "template.py",
        BACKEND_ROOT / "_v21_r9_preview_pdf.py",
    ]
    compile_ok = True
    for p in targets:
        try:
            py_compile.compile(str(p), doraise=True)
        except py_compile.PyCompileError as e:
            compile_ok = False
            print(f"  py_compile 失败：{p}: {e}", file=sys.stderr)
    _assert(compile_ok, "改动文件 py_compile 全部通过")

    import_ok = True
    err = ""
    try:
        from api.routes import template as tr
        from services import pdf_renderer as pr
        from services import resume_generation_service as rgs
        if tr.router is None or pr is None or rgs is None:
            raise RuntimeError("import 返回 None")
    except Exception as ex:
        import_ok = False
        err = repr(ex)
    _assert(import_ok, "template/pdf_renderer/resume_generation_service 模块 import 成功", err)


# ═══════════════════════════════════════════════════════════════════════════
# 入口
# ═══════════════════════════════════════════════════════════════════════════
def main() -> int:
    doc, docx_path, pdf_bytes, pdf_path, (_dw, _pw) = _produce_all()
    test_pdf_authentic(pdf_bytes, pdf_path, "pm_template")
    test_triplet_consistency(doc, docx_path, pdf_path)
    test_download_endpoint(pdf_bytes, pdf_path, docx_path)
    test_response_contract_and_compile()

    print()
    print("=" * 64)
    print(f"V2.1.0 R9 三端一致性/PDF 渲染：PASS={PASS_COUNT} FAIL={len(FAILURES)}")
    print("=" * 64)
    if FAILURES:
        for f in FAILURES:
            print(f"  - {f}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    code = main()
    print(f"PASS={PASS_COUNT} FAIL={len(FAILURES)}")
    sys.exit(code)
