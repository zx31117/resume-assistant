"""V1.2 DOCX Writer（简化版，不用文本框/坐标）。

核心职责（只处理 Paragraph 级操作，不理解章节业务）：
  - clone_paragraph：深拷贝原型段 XML，保留 style + 所有 Run 的 rPr
  - fill_placeholders：只改 w:t 文本，不动 rPr（V1.2 简化版，不做 token 拆分）
  - insert_after：在某段之后插入新段
  - remove_paragraph：从文档中删除段落

不做：
  - 文本框 / shape / SmartArt 读写
  - 章节业务识别（由 TemplateRenderer 负责理解 section_spec）
  - Run 级 token 拆分（留 V2；V1.2 靠模板编写规范保证占位符独占 Run）
"""
import copy
import re
import os
import json
from typing import Optional

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from models.template_schema import TemplateSpec


# ── 基础 Paragraph 操作 ────────────────────────────────────────── #

def clone_paragraph(proto_p: Paragraph) -> Paragraph:
    """深拷贝段落 XML，保留 style + 所有 Run 的 rPr。返回的新段未插入文档。

    - proto_p._p 整个 w:p 深拷贝
    - Run 的 rPr（字体/字号/加粗）完整保留
    - w:t 的 text 不做任何修改（调用方之后会调用 fill_placeholders 替换占位符）
    """
    new_p_elem = copy.deepcopy(proto_p._p)
    return Paragraph(new_p_elem, proto_p._parent)


def insert_after(ref_p: Paragraph, new_p: Paragraph) -> Paragraph:
    """在 ref_p 之后插入 new_p（ref_p 需已在文档里）。返回 new_p。"""
    ref_p._p.addnext(new_p._p)
    return new_p


def remove_paragraph(p: Paragraph) -> None:
    """从文档中移除该段落。"""
    p._p.getparent().remove(p._p)


def find_paragraphs_by_style(doc: Document, style_name: str) -> list[Paragraph]:
    """按段落 style.name 精确匹配，返回所有命中段落（渲染器按 style 定位，不靠文本）。"""
    return [p for p in doc.paragraphs if p.style.name == style_name]


def find_next_paragraph(p: Paragraph) -> Optional[Paragraph]:
    """取 p 的下一个兄弟段落（非 w:p 的跳过）。找不到返回 None。"""
    node = p._p.getnext()
    while node is not None and node.tag != qn('w:p'):
        node = node.getnext()
    if node is None:
        return None
    return Paragraph(node, p._parent)


# ── V1.2 简化版占位符替换（不做 token 拆分） ──────────────────── #
#
# 前提：V1.2 模板编写规范 R5：一个 {{xxx}} 占位符独占一个 Run
#
# 因此 fill_placeholders 只做两件事：
#   1. 遍历段落里所有 Run，如果某个 Run 的 w:t 完全就是 {{xxx}}（前后只有空白）
#      → 用 ctx 里对应的值替换
#   2. 如果某个 Run 的 w:t 包含多个 {{xxx}}（模板不规范，兜底）→ 字符串级 replace 拼接
#

_PH_SINGLE_RE = re.compile(r'^\s*\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}\s*$')
_PH_MULTI_RE = re.compile(r'\{\{\s*([a-zA-Z0-9_.]+)\s*\}\}')

def _resolve_dotpath(ctx: dict, dotpath: str):
    """按点路径解析 ctx 里的值。

    ctx 可以是嵌套 dict，或者 ResumeDocument 对象（用 getattr/字典兼容）。
    """
    parts = dotpath.split(".")
    cur: object = ctx
    for part in parts:
        if isinstance(cur, dict):
            if part not in cur:
                return None
            cur = cur[part]
        else:
            if not hasattr(cur, part):
                return None
            cur = getattr(cur, part)
            if callable(cur):
                return None
    if cur is None:
        return ""
    if isinstance(cur, list):
        # 列表字段（如 skill.items）固定按「顿号分隔」字符串化（V1.2 代码内固定，不做 join filter）
        return "、".join(str(x) for x in cur if x is not None)
    return str(cur)


def fill_placeholders(p: Paragraph, ctx: dict) -> bool:
    """替换段落中所有 {{xxx}} 占位符。

    返回 True 表示段落里任意占位符解析后为空（即该段需要在渲染器逻辑中删除）。
    """
    any_empty = False

    for run in p.runs:
        text = run.text or ""
        if "{{" not in text:
            continue

        m = _PH_SINGLE_RE.match(text)
        if m:
            # Case A：整个 Run 就是一个占位符 —— 直接替换整个 w:t
            dotpath = m.group(1)
            val = _resolve_dotpath(ctx, dotpath)
            if not val:
                any_empty = True
                run.text = ""
            else:
                run.text = val
            continue

        # Case B：占位符和其他字符混在一个 Run（模板不规范的兜底）
        # 字符串级 replace，rPr 保持不变
        def sub_fn(match):
            dotpath = match.group(1)
            val = _resolve_dotpath(ctx, dotpath)
            if not val:
                nonlocal any_empty
                any_empty = True
                return ""
            return str(val)
        run.text = _PH_MULTI_RE.sub(sub_fn, text)

    return any_empty


# ── 文档级：加载模板资产（docx + json 成对） ───────────────────── #

def load_template_assets(template_id: str, backend_root: str) -> tuple[Document, TemplateSpec]:
    """按 template_id 从 config/template_mapping.json 找 docx + json 路径并加载。

    返回 (python-docx Document 实例（内存副本）, TemplateSpec 对象)
    """
    mapping_path = os.path.join(backend_root, "config", "template_mapping.json")
    with open(mapping_path, "r", encoding="utf-8") as f:
        mapping = json.load(f)
    if template_id not in mapping:
        raise KeyError(f"template_id={template_id!r} 未在 config/template_mapping.json 注册")
    entry = mapping[template_id]

    docx_rel = entry["docx"].replace("/", os.sep)
    json_rel = entry["json"].replace("/", os.sep)
    docx_path = os.path.join(backend_root, docx_rel)
    json_path = os.path.join(backend_root, json_rel)

    if not os.path.exists(docx_path):
        raise FileNotFoundError(
            f"模板 DOCX 不存在: {docx_path}. "
            f"请先在 backend/ 目录下运行: python templates/_build_templates.py"
        )

    doc = Document(docx_path)

    with open(json_path, "r", encoding="utf-8") as f:
        spec = TemplateSpec(**json.load(f))

    return doc, spec


# ── 异常类（渲染时使用） ────────────────────────────────────────── #

class TemplateError(Exception):
    """模板结构错误（比如 section_spec.title_style 在 docx 里找不到对应样式段落）。"""
    pass
