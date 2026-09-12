"""V2.1.0 H6（PLAN §18.2）确定性虚构 fixture 生成器（零 API Key、零隐私、可审计）。

产出：
  build_fixture_pdf(artifact_id, bullet_fact_refs=None) -> (pdf_bytes, anchors)
  build_fixture_docx() -> docx_bytes
  load_resume_doc()

确定性设计（同一 requirements.txt 版本 + 仓库内同一字体下字节级可复现）：
- PDF：reportlab Canvas(..., invariant=1) —— 关闭时间戳/随机 ID，输出与日期无关；
  中文用仓库随附固定字体 backend/templates/fonts/NotoSansSC-Regular.ttf（OFL，R17a 已固化）。
- DOCX：python-docx 构建后做「确定性重打包」——固定 docProps/core.xml 的
  dcterms:created/modified 时间 + 全部 ZIP 条目时间戳置为固定值、ZIP_STORED 不压缩，
  使输出字节不依赖当前时钟与 zlib 压缩实现。

安全边界：
- 内容唯一来源 resume_doc.json / jd.txt（全虚构；_meta.privacy_check 已登记）；
- 所有路径由 __file__ 解析（仓库相对），不出现本机绝对路径；
- 不读取任何真实 runtime、数据库、网络或 Key；导入/运行不产生任何写入副作用。
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

FIXTURE_DIR = Path(__file__).resolve().parent
BACKEND_ROOT = FIXTURE_DIR.parent
FONT_PATH = BACKEND_ROOT / "templates" / "fonts" / "NotoSansSC-Regular.ttf"
RESUME_DOC_JSON = FIXTURE_DIR / "resume_doc.json"

# ── A4 布局参数（与 pdf_renderer 同约定：pt、y 自底向上，锚点坐标在 A4 页内）──
A4_W, A4_H = 595.28, 842.0
CM = 28.3465
MARGIN_LEFT = 1.13 * CM
MARGIN_RIGHT = 1.09 * CM
MARGIN_TOP = 0.92 * CM
MARGIN_BOTTOM = 1.39 * CM
CONTENT_W = A4_W - MARGIN_LEFT - MARGIN_RIGHT
_EM_ASCENT = 0.88
_EM_DESCENT = 0.12

BULLET = "● "
SZ_NAME = 18.0
SZ_BODY = 10.6
SZ_TITLE = 12.0
SZ_CONTACT = 10.0
LEAD_BODY = 17.0
LEAD_TITLE = 16.0

_SECTION_TITLES = {
    "education": "教育背景",
    "work": "实习经历",
    "project": "项目经历",
    "skills": "技能专长",
    "awards": "荣誉奖项",
    "summary": "自我评价",
}


def load_resume_doc() -> dict:
    return json.loads(RESUME_DOC_JSON.read_text(encoding="utf-8"))


def _doc_sections(doc: dict) -> list[tuple[str, list]]:
    """按模板渲染顺序收集「非空即可见」的章节（education/work/project/skills/awards/summary）。"""
    out: list[tuple[str, list]] = []
    prof = doc.get("profile") or {}
    for sec, key in (
        ("education", "education"), ("work", "work"),
        ("project", "projects"), ("skills", "skills"), ("awards", "awards"),
    ):
        items = doc.get(key) or []
        if sec == "skills":
            visible = [g for g in items if g.get("category") and g.get("items")]
        else:
            visible = [x for x in items if x]
        if visible:
            out.append((sec, visible))
    summary_lines = [ln.strip() for ln in (prof.get("summary") or "").split("\n") if ln.strip()]
    if summary_lines:
        out.append(("summary", summary_lines))
    return out


def _approx_w(text: str, size: float = SZ_BODY) -> float:
    """CJK 为主的估算字符宽（仅用于标题右对齐排版，不影响锚点记录）。"""
    return len(text) * size * 0.95


# ═══════════════════════════════════════════════════════════════════════════
# PDF：reportlab invariant 两页（技能专长前强制换页 → page_index 0/1 均有锚点）
# ═══════════════════════════════════════════════════════════════════════════
def build_fixture_pdf(artifact_id: str = "", bullet_fact_refs: dict | None = None):
    """渲染固定虚构简历为两页 PDF。

    返回 (pdf_bytes, anchors)。anchors 与 pdf_renderer.PreviewAnchor 同约定：
      {artifact_id, page_index, x0,y0,x1,y1, content_item_id, bullet_index, text, fact_refs}
    坐标 pt、y 自底向上、A4 页内；fact_refs 逐 bullet 取自 bullet_fact_refs，无映射一律 []。
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas as rl_canvas

    if not FONT_PATH.is_file():
        raise RuntimeError(
            f"H6 fixture 字体缺失：{FONT_PATH.name}（应随仓库 backend/templates/fonts/ 提供）"
        )
    pdfmetrics.registerFont(TTFont("H6Noto", str(FONT_PATH)))

    doc = load_resume_doc()
    refs_map = {k: list(v or []) for k, v in (bullet_fact_refs or {}).items()}

    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle("H6-fict-resume")
    anchors: list[dict] = []
    page_no = 0
    # b：当前文本 baseline（自底向上），初始在内容顶缘。
    b = A4_H - MARGIN_TOP

    def new_page() -> None:
        nonlocal page_no, b
        c.showPage()
        page_no += 1
        b = A4_H - MARGIN_TOP

    def ensure_room(need_pt: float = LEAD_BODY) -> None:
        nonlocal b
        if b - need_pt < MARGIN_BOTTOM:
            new_page()

    def emit(text: str, x: float, y: float, size: float, bold: bool = False) -> None:
        if not text:
            return
        t = c.beginText(x, y)
        t.setFont("H6Noto", size)
        if bold:
            t.setTextRenderMode(2)  # fill+stroke 模拟粗体（同 pdf_renderer）
            c.setLineWidth(0.6)
        t.textOut(text)
        c.drawText(t)
        if bold:
            c.setLineWidth(1.0)

    def flow_lines(text: str, indent: float, size: float):
        """把 text 按真实字体宽度折行渲染；返回 [(行文本, 行宽 pt), ...]。"""
        nonlocal b
        from reportlab.pdfbase import pdfmetrics as _pm
        avail = CONTENT_W - indent
        lines: list[tuple[str, float]] = []
        cur = ""
        for ch in text:
            if cur and _pm.stringWidth(cur + ch, "H6Noto", size) > avail:
                lines.append((cur, _pm.stringWidth(cur, "H6Noto", size)))
                cur = ch
            else:
                cur += ch
        if cur or not lines:
            lines.append((cur, _pm.stringWidth(cur, "H6Noto", size)))
        for i, (ln, _w) in enumerate(lines):
            if indent and i == 0:
                emit(BULLET, MARGIN_LEFT, b, SZ_BODY)
            emit(ln, MARGIN_LEFT + indent, b, size)
            b -= LEAD_BODY if size <= SZ_BODY else (SZ_NAME + 4.0)
        return lines

    def draw_section_title(title: str) -> None:
        nonlocal b
        emit(title, MARGIN_LEFT, b, SZ_TITLE, bold=True)
        b -= 4.0
        c.setLineWidth(0.75)
        c.line(MARGIN_LEFT, b, A4_W - MARGIN_RIGHT, b)
        b -= LEAD_TITLE

    def draw_item_title(item: dict, sec: str) -> None:
        nonlocal b
        st, en = item.get("start_time") or "", item.get("end_time") or ""
        period = f"{st}-{en}" if (st and en) else (st or en or "")
        if sec == "education":
            left, center, right = period, item.get("school") or "", (
                f"{item.get('major')}（{item.get('degree')}）" if item.get("degree")
                else (item.get("major") or "")
            )
        elif sec == "work":
            left, center, right = period, item.get("company") or "", item.get("role") or ""
        else:
            left, center, right = period, item.get("name") or "", item.get("role") or ""
        if left:
            emit(left, MARGIN_LEFT, b, SZ_BODY, bold=True)
        if center:
            emit(center, MARGIN_LEFT + (CONTENT_W - _approx_w(center)) / 2.0, b, SZ_BODY, bold=True)
        if right:
            emit(right, A4_W - MARGIN_RIGHT - _approx_w(right), b, SZ_BODY, bold=True)
        b -= LEAD_BODY

    def draw_anchor(text: str, content_item_id, bullet_index: int,
                    fact_refs, indent: float, size: float = SZ_BODY) -> None:
        nonlocal b
        first_b = b
        lines = flow_lines(text, indent, size)  # flow_lines 已把 b 推进 len(lines) 行
        max_w = max((w for _ln, w in lines), default=0.0)
        last_b = first_b - LEAD_BODY * (len(lines) - 1)  # 末行 baseline（自底向上坐标）
        anchors.append({
            "artifact_id": artifact_id,
            "page_index": page_no,
            "x0": round(MARGIN_LEFT, 2),
            "y0": round(last_b - size * _EM_DESCENT, 2),
            "x1": round(min(MARGIN_LEFT + CONTENT_W, MARGIN_LEFT + indent + max_w) + 2.0, 2),
            "y1": round(first_b + size * _EM_ASCENT, 2),
            "content_item_id": content_item_id,
            "bullet_index": bullet_index,
            "text": text,
            "fact_refs": list(fact_refs or []),
        })

    def bullet_refs(eid: str, j: int) -> list[str]:
        prefs = refs_map.get(eid) if eid else None
        if prefs and 0 <= j < len(prefs):
            return list(prefs[j] or [])
        return []

    # ── profile 头部 ──
    prof = doc.get("profile") or {}
    if prof.get("name"):
        flow_lines(prof["name"], 0.0, SZ_NAME)
    if prof.get("target_position"):
        flow_lines(f"求职意向：{prof['target_position']}", 0.0, SZ_CONTACT)
    contacts = [p for p in (
        f"电话：{prof.get('phone')}" if prof.get("phone") else "",
        f"邮箱：{prof.get('email')}" if prof.get("email") else "",
        f"所在地：{prof.get('location')}" if prof.get("location") else "",
    ) if p]
    if len(contacts) == 3:
        flow_lines(" 丨 ".join(contacts), 0.0, SZ_CONTACT)
    b -= 6.0

    # ── 章节（education/work/project 第 1 页开始；skills 起强制新页）──
    for sec, items in _doc_sections(doc):
        if sec == "skills" and page_no == 0:
            new_page()  # 保证两页几何：第 2 页从技能专长开始
        draw_section_title(_SECTION_TITLES[sec])
        if sec in ("education", "work", "project"):
            for item in items:
                eid = item.get("experience_id") or ""
                draw_item_title(item, sec)
                if sec == "education":
                    bi = 0
                    for field in ("description", "gpa"):
                        val = (item.get(field) or "").strip()
                        if val:
                            ensure_room()
                            draw_anchor(val, eid, bi, [], indent=14.0)
                            bi += 1
                    for j, bl in enumerate(item.get("bullets") or []):
                        if bl.strip():
                            ensure_room()
                            draw_anchor(bl.strip(), eid, bi, bullet_refs(eid, j), indent=14.0)
                            bi += 1
                else:
                    for j, bl in enumerate(item.get("bullets") or []):
                        if bl.strip():
                            ensure_room()
                            draw_anchor(bl.strip(), eid, j, bullet_refs(eid, j), indent=14.0)
        elif sec == "skills":
            for i, g in enumerate(items):
                joined = "、".join(g.get("items") or [])
                ensure_room()
                draw_anchor(f"{g.get('category')}：{joined}", f"skills:{i}", i, [], indent=0.0)
        elif sec == "awards":
            for i, a in enumerate(items):
                ensure_room()
                draw_anchor(str(a), "awards", i, [], indent=14.0)
        elif sec == "summary":
            for i, ln in enumerate(items):
                ensure_room()
                draw_anchor(str(ln), "summary", i, [], indent=14.0)

    c.showPage()
    c.save()
    return buf.getvalue(), anchors


# ═══════════════════════════════════════════════════════════════════════════
# DOCX：python-docx + 确定性重打包
# ═══════════════════════════════════════════════════════════════════════════
_FIXED_TS = (2020, 1, 1, 0, 0, 0)
_FIXED_ISO = "2020-01-01T00:00:00Z"


def _deterministic_docx(raw: bytes) -> bytes:
    """把 python-docx 输出重打包为字节确定形式（时间戳固定 + ZIP_STORED）。"""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)",
                    rf"\g<1>{_FIXED_ISO}\g<2>", text,
                )
                text = re.sub(
                    r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rf"\g<1>{_FIXED_ISO}\g<2>", text,
                )
                data = text.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=_FIXED_TS)
            zi.compress_type = zipfile.ZIP_STORED
            zi.external_attr = 0
            z.writestr(zi, data)
    return out.getvalue()


def build_fixture_docx() -> bytes:
    """用 python-docx 构建固定虚构简历 DOCX（确定性字节）。"""
    import docx as docx_mod

    doc = load_resume_doc()
    prof = doc.get("profile") or {}
    d = docx_mod.Document()
    d.add_heading(prof.get("name") or "", level=0)
    d.add_paragraph("全职虚构简历样本（V2.1.0 H6 fixture，仅供专项测试）")
    edu_items = doc.get("education") or []
    for ei, item in enumerate(edu_items):
        if ei == 0:
            d.add_heading("教育背景", level=1)
            d.add_paragraph(
                f"{item.get('school')} · {item.get('major')}（{item.get('degree')}） "
                f"{item.get('start_time')}-{item.get('end_time')}"
            )
            if item.get("description"):
                d.add_paragraph(item["description"], style="List Bullet")
            if item.get("gpa"):
                d.add_paragraph(item["gpa"], style="List Bullet")
    for sec, title, key in (("work", "实习经历", "work"), ("project", "项目经历", "projects")):
        d.add_heading(title, level=1)
        for item in doc.get(key) or []:
            if sec == "work":
                head = f"{item.get('company')} · {item.get('role')} {item.get('start_time')}-{item.get('end_time')}"
            else:
                head = f"{item.get('name')} · {item.get('role')} {item.get('start_time')}-{item.get('end_time')}"
            d.add_paragraph(head)
            for bl in item.get("bullets") or []:
                if bl.strip():
                    d.add_paragraph(bl.strip(), style="List Bullet")
    d.add_heading("技能专长", level=1)
    for g in doc.get("skills") or []:
        d.add_paragraph(f"{g.get('category')}：{'、'.join(g.get('items') or [])}")
    buf = io.BytesIO()
    d.save(buf)
    return _deterministic_docx(buf.getvalue())


if __name__ == "__main__":
    import hashlib
    import pdfplumber

    doc = load_resume_doc()
    pdf, anchors = build_fixture_pdf(
        artifact_id="H6-ARTIFACT-TOKEN", bullet_fact_refs=doc.get("bullet_fact_refs") or {},
    )
    docx = build_fixture_docx()
    with pdfplumber.open(io.BytesIO(pdf)) as p:
        print("pages:", len(p.pages))
        for i, pg in enumerate(p.pages):
            print(f"--- page {i} head ---")
            print("\n".join((pg.extract_text() or "").splitlines()[:4]))
    print("pdf sha256:", hashlib.sha256(pdf).hexdigest(), "len", len(pdf))
    print("docx sha256:", hashlib.sha256(docx).hexdigest(), "len", len(docx))
    print("anchors:", len(anchors))
    for a in anchors:
        print(
            " ", a["page_index"], a["content_item_id"], a["bullet_index"],
            a["fact_refs"], a["x0"], a["y0"], a["x1"], a["y1"], a["text"][:20],
        )
